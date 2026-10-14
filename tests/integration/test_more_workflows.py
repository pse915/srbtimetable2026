import base64
from datetime import date
import io
import json
import pandas as pd
from openpyxl import load_workbook
from tests.fixtures import make_application,login,click,command

def test_calendar_weekday_and_week_navigation():
    app,_=make_application();login(app)
    calendar=next(k for k,v in app.port.controls.items() if v['node']['kind']=='calendar')
    result=command(app,'control',{'control':calendar,'values':{calendar:'2026-10-13'}})
    assert result['response']['success']
    assert app.port.session_state['_view_daily_date_selected']==date(2026,10,13)
    result=command(app,'control',{'control':calendar,'values':{calendar:'2026-10-17'}})
    assert not result['response']['success']

def test_dnd_preview_requires_approval_and_reuses_validator():
    app,sheets=make_application();login(app);command(app,'navigate',{'tab':'시간표 맞교환 & 변경 추천'})
    table=next(k for k,v in app.port.controls.items() if str(k).startswith('exchange_weekly_matrix__'))
    records=app.port.controls[table]['node']['table']['rows']
    source=next(i for i,r in enumerate(records) if r['교사명']=='테스트가')
    target=next(i for i,r in enumerate(records) if r['교사명']=='테스트나')
    preview=command(app,'preview_swap',{'control':table,'source':[source,'월1'],'target':[target,'화2']})
    assert preview['response']['success']
    assert sheets['맞교환'].writes==0
    token=preview['swapPreview']['token']
    result=command(app,'apply_swap_preview',{'token':token})
    assert result['response']['success'] and sheets['맞교환'].writes==1
    result=command(app,'apply_swap_preview',{'token':token})
    assert not result['response']['success'] and sheets['맞교환'].writes==1

def test_undo_redo_budget_data_and_history_cap():
    app,_=make_application();login(app);ns=app.program.namespace
    for i in range(7):
        app.port.session_state.duties=pd.DataFrame([{'교사명':'테스트가','일자':'2026-10-12','교시':i}])
        ns['push_history'](f'변경{i}')
    assert len(app.port.session_state.history)==5
    assert ns['undo']() and app.port.session_state.duties.iloc[0]['교시']==5
    assert ns['redo']() and app.port.session_state.duties.iloc[0]['교시']==6

def test_salary_editor_and_type_rate_default_change():
    app,_=make_application();login(app);command(app,'navigate',{'tab':'교무호봉획정'});click(app,'+ 경력 추가')
    control=next(k for k,v in app.port.controls.items() if str(k).startswith('salary_career_editor'))
    career_type=app.program.namespace['SALARY_CAREER_TYPES'][-1]
    rows=[{'경력 종류':career_type,'세부 근무처/직위':'테스트','시작일':'2022-03-01','종료일':'2023-02-28','종료일 산입':True,'환산율 (%)':100}]
    command(app,'control',{'control':control,'values':{control:rows}})
    assert app.port.session_state.salary_careers.iloc[0]['환산율 (%)']==40
    click(app,'📊 호봉 계산 및 조서 생성')
    assert app.port.session_state.salary_result['total_m']==4 and app.port.session_state.salary_result['total_d']==24

def test_weekly_excel_output_content_and_print_settings():
    app,_=make_application();login(app)
    raw=app.program.call('build_weekly_schedule_excel_bytes',date(2026,10,12))
    ws=load_workbook(io.BytesIO(raw)).active
    assert ws.freeze_panes=='C5' and ws.page_setup.orientation=='landscape'
    assert ws['B5'].value=='테스트가' and ws['C5'].value=='국어' and ws['C6'].value=='1-1'
    raw=app.program.call('build_weekly_class_schedule_excel_bytes',date(2026,10,12))
    ws=load_workbook(io.BytesIO(raw)).active
    assert ws['B5'].value=='1-1' and ws['C6'].value=='테스트가'

def test_gemini_payload_and_ai_result_use_original_config(monkeypatch):
    app,_=make_application();login(app);ns=app.program.namespace
    app.port.secrets['GEMINI_API_KEY']='TEST-GEMINI-KEY'
    from backend.services.application import UploadedFile
    data={'name':'분석이름','baseSalary':8,'isSabom':False,'careers':[{'type':ns['SALARY_CAREER_TYPES'][-1],'start':'2022-03-01','end':'2023-02-28','rate':100,'inc':True}]}
    seen={}
    class Response:
        status_code=200
        def json(self):return {'candidates':[{'content':{'parts':[{'text':json.dumps(data,ensure_ascii=False)}]}}]}
    def post(url,**kwargs):seen.update(kwargs);return Response()
    monkeypatch.setattr(ns['requests'],'post',post)
    result=ns['_salary_analyze_uploaded_file'](UploadedFile(b'%PDF-1.7 test','test.pdf','application/pdf'))
    ns['_salary_init_state']();ns['_salary_apply_ai_result'](result)
    assert seen['params']=={'key':'TEST-GEMINI-KEY'}
    assert base64.b64decode(seen['json']['contents'][0]['parts'][0]['inlineData']['data'])==b'%PDF-1.7 test'
    assert app.port.session_state.salary_name=='분석이름' and app.port.session_state.salary_careers.iloc[0]['환산율 (%)']==40

def test_neis_info_200_fallback_and_non_200_error_propagation():
    from backend.bridge.port import ReactPort
    from backend.domain.program import LegacyProgram
    original=LegacyProgram(ReactPort()).namespace
    calls=[]
    original['neis_find_school']=lambda *args:{'ATPT_OFCDC_SC_CODE':'R10','SD_SCHUL_CODE':'8771121'}
    def fetch(key,school,**kw):
        calls.append(kw)
        if 'ymd' not in kw:raise RuntimeError('NEIS INFO-200')
        return {'SchoolSchedule':[{'row':[{'AA_YMD':kw['ymd'],'EVENT_NM':'휴업일'}]}]}
    original['_neis_schedule_request']=fetch
    df=original['neis_fetch_schedule']('key','school','2026-10-12','2026-10-16')
    assert len(calls)==6 and len(df)==5 and df['비수업일'].all()
    original['neis_fetch_schedule'].clear()
    def fail(*args,**kw):raise RuntimeError('HTTP 403')
    original['_neis_schedule_request']=fail
    import pytest
    with pytest.raises(RuntimeError,match='403'):original['neis_fetch_schedule']('key','school','2026-10-12','2026-10-16')

def test_ai_updates_previously_rendered_widget_values():
    app,_=make_application();login(app);command(app,'navigate',{'tab':'교무호봉획정'})
    ns=app.program.namespace
    ns['_salary_apply_ai_result']({'name':'AI반영이름','isSabom':False,'degrees':[],'careers':[]})
    app._render()
    assert app.port.session_state.salary_name=='AI반영이름'
    assert app.port.session_state.salary_name_input=='AI반영이름'
    assert app.port.session_state.salary_sabom_checkbox is False

def test_salary_remains_accessible_when_timetable_read_fails():
    app,sheets=make_application();sheets['시간표'].values=[]
    login(app)
    result=command(app,'navigate',{'tab':'교무호봉획정'})
    assert result['response']['success']
    assert any(s['node']['label']=='성명' for s in app.port.controls.values())

def test_logout_clears_private_drafts_files_and_request_receipts():
    app,_=make_application();login(app);command(app,'navigate',{'tab':'교무호봉획정'});click(app,'📊 호봉 계산 및 조서 생성')
    old_ids=set(app.responses)
    click(app,'로그아웃')
    assert not app.port.session_state.get('logged_in')
    assert not old_ids.intersection(app.responses)
    assert not app.uploads and not app.port.session_state.get('salary_result')
